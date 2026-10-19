#!/usr/bin/env python3
"""Extract invoice fields from sample PDFs → Excel + CSV. No paid API."""
from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
DEFAULT_SAMPLES = ROOT / "samples"
DEFAULT_OUT_XLSX = ROOT / "output" / "invoices.xlsx"
DEFAULT_OUT_CSV = ROOT / "output" / "invoices.csv"

FIELD_PATTERNS = {
    "invoice_no": re.compile(r"Invoice\s*#\s*:\s*(.+)", re.I),
    "date": re.compile(r"Invoice\s*Date\s*:\s*(\d{4}-\d{2}-\d{2})", re.I),
    "due": re.compile(r"Due\s*Date\s*:\s*(\d{4}-\d{2}-\d{2})", re.I),
    "currency": re.compile(r"Currency\s*:\s*([A-Z]{3})", re.I),
    "total": re.compile(r"(?m)^TOTAL\s*:\s*([\d.,]+)\s*([A-Z]{3})?", re.I),
    "subtotal": re.compile(r"(?m)^Subtotal\s*:\s*([\d.,]+)", re.I),
    "vat": re.compile(r"(?m)^VAT\s*:\s*([\d.,]+)", re.I),
}


def pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def parse_invoice(path: Path) -> dict:
    text = pdf_text(path)
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    vendor = lines[1] if len(lines) > 1 else ""
    # Skip title "INVOICE"
    if lines and lines[0].upper() == "INVOICE" and len(lines) > 1:
        vendor = lines[1]
    row = {
        "file": path.name,
        "vendor": vendor,
        "invoice_no": "",
        "date": "",
        "due": "",
        "currency": "",
        "subtotal": "",
        "vat": "",
        "total": "",
    }
    for key, pat in FIELD_PATTERNS.items():
        m = pat.search(text)
        if not m:
            continue
        if key == "total":
            row["total"] = m.group(1).replace(",", "")
            if m.group(2):
                row["currency"] = m.group(2)
        else:
            row[key] = m.group(1).strip().replace(",", "")
    return row


def write_excel(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    headers = ["file", "vendor", "invoice_no", "date", "due", "currency", "subtotal", "vat", "total"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for r_i, row in enumerate(rows, 2):
        for c_i, h in enumerate(headers, 1):
            val = row.get(h, "")
            if h in ("subtotal", "vat", "total") and val != "":
                try:
                    val = float(val)
                except ValueError:
                    pass
            cell = ws.cell(r_i, c_i, val)
            cell.border = thin
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col) + 2
        ws.column_dimensions[col[0].column_letter].width = min(width, 36)
    # Summary sheet
    ws2 = wb.create_sheet("Summary")
    ws2["A1"] = "Currency"
    ws2["B1"] = "Invoice count"
    ws2["C1"] = "Total sum"
    for cell in (ws2["A1"], ws2["B1"], ws2["C1"]):
        cell.font = Font(bold=True)
    by_cur: dict[str, list[float]] = {}
    for row in rows:
        cur = row.get("currency") or "?"
        try:
            tot = float(row.get("total") or 0)
        except ValueError:
            tot = 0.0
        by_cur.setdefault(cur, []).append(tot)
    for i, (cur, vals) in enumerate(sorted(by_cur.items()), 2):
        ws2.cell(i, 1, cur)
        ws2.cell(i, 2, len(vals))
        ws2.cell(i, 3, round(sum(vals), 2))
    wb.save(path)


def write_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = ["file", "vendor", "invoice_no", "date", "due", "currency", "subtotal", "vat", "total"]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        for row in rows:
            w.writerow({h: row.get(h, "") for h in headers})


def main() -> None:
    ap = argparse.ArgumentParser(description="PDF invoices → Excel")
    ap.add_argument("--samples", type=Path, default=DEFAULT_SAMPLES)
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_OUT_XLSX)
    ap.add_argument("--csv", type=Path, default=DEFAULT_OUT_CSV)
    args = ap.parse_args()
    pdfs = sorted(args.samples.glob("*.pdf"))
    if not pdfs:
        raise SystemExit(f"No PDFs in {args.samples}")
    rows = [parse_invoice(p) for p in pdfs]
    write_excel(rows, args.xlsx)
    write_csv(rows, args.csv)
    print(f"Parsed {len(rows)} invoices")
    for r in rows:
        print(f"  {r['invoice_no']:12} {r['vendor'][:28]:28} {r['total']:>8} {r['currency']}")
    print(f"Excel → {args.xlsx}")
    print(f"CSV   → {args.csv}")


if __name__ == "__main__":
    main()
